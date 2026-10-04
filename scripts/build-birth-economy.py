#!/usr/bin/env python3
"""v123 · Construye public/birth-economy.json ("La economía cuando naciste").

Entradas (descargadas de fuentes oficiales; ver README):
  series.json  -> Datos Argentina / Banco Mundial (IPC histórico INDEC, IPC San Luis, IPC Nacional,
                  tipo de cambio BCRA en $ equivalentes, SMVM, RIPTE, PIB 1935-62, precios promedio INDEC)
  uscpi.csv    -> FRED CPIAUCNS (IPC-U EE.UU., mensual desde 1913)
  people/*.json, wp.json -> presidentes, ministros y fotos (Wikipedia / Wikimedia Commons)
  goods.json, staples.json -> precios de auto 0 km, m² en CABA, pan, asado, leche (con fuente)
  public/fx-daily.json, public/macro-history.json -> dólar libre diario 1991+ y anual histórico

Todos los montos se guardan en "pesos actuales equivalentes" (moneda vigente desde 1992);
el navegador los convierte a la moneda de cada época y a valores de hoy con el IPC encadenado.
Uso: python3 scripts/build-birth-economy.py <carpeta-de-insumos>
"""
import json, sys, csv, datetime, os

W = sys.argv[1] if len(sys.argv) > 1 else '.'
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
S = json.load(open(os.path.join(W, 'series.json')))
ym = lambda d: d[:7]
def mp(rows): return {ym(a): b for a, b in rows if b is not None}

START = '1943-01'
def months(a, b):
    y, m = map(int, a.split('-')); out = []
    while f'{y}-{m:02d}' <= b:
        out.append(f'{y}-{m:02d}'); m += 1
        if m > 12: m = 1; y += 1
    return out

# ---------- IPC encadenado (misma regla que el gráfico de inflación del sitio) ----------
H, SL, N = mp(S['cpiHist']), mp(S['cpiSL']), mp(S['cpiNat'])
last_cpi = max(N)
M = months(START, last_cpi)
idx = {START: 1.0}
for prev, cur in zip(M, M[1:]):
    y = int(cur[:4])
    src = H if y <= 2006 else SL if y <= 2016 else N
    if prev in src and cur in src:
        idx[cur] = idx[prev] * src[cur] / src[prev]
    else:
        raise SystemExit(f'IPC sin dato en {cur}')
base = idx[last_cpi]
cpi = {k: v / base for k, v in idx.items()}           # 1.0 = último mes publicado

# ---------- Tipo de cambio oficial BCRA ($ equivalentes) + extensión con fx-daily ----------
def expand(s):
    out = {}; dt = datetime.date.fromisoformat(s['start'])
    for d, v in zip(s['d'], s['v']):
        dt = dt + datetime.timedelta(days=d or 0)
        if v is not None: out[dt.isoformat()] = v
    return out
fxd = json.load(open(os.path.join(ROOT, 'public', 'fx-daily.json')))
def monthly_avg(daily):
    acc = {}
    for d, v in daily.items(): acc.setdefault(d[:7], []).append(v)
    return {k: sum(v) / len(v) for k, v in acc.items()}
off_m = monthly_avg(expand(fxd['official'])); free_m = monthly_avg(expand(fxd['free']))
fx = mp(S['fxEq'])
for k, v in off_m.items():
    if k > max(fx): fx[k] = v                               # meses posteriores a la serie BCRA
hist = json.load(open(os.path.join(ROOT, 'public', 'macro-history.json')))
fx_free_annual = hist['marketsHistory']['exchangeHistorical']['free']['nominal']   # $ actuales, anual

# ---------- Salarios ----------
# v129: hasta dic-1993, salario medio industrial nominal (data/salarios.xlsx, serie mensual 1940-1993 con su
# moneda por fila); desde jul-1994, RIPTE. Ene-jun 1994 se completa interpolando en términos reales (IPC).
import openpyxl, re
CURNAME = {'Peso moneda nacional (m$n)': 1e-13, 'Peso Ley 18.188 ($ ley)': 1e-11, 'Peso argentino ($a)': 1e-7, 'Austral (₳)': 1e-4, 'Peso ($)': 1.0}
today = datetime.date.today().strftime('%Y-%m')
ripte = mp(S['ripte'])
wage, wage_src = {}, {}
_wb = openpyxl.load_workbook(os.path.join(ROOT, 'data', 'salarios.xlsx'), data_only=True)
for r in list(_wb['Serie mensual 1940-1993'].iter_rows(values_only=True))[1:]:
    if not r[1] or r[3] is None: continue
    k = f"{int(r[1])}-{int(r[2]):02d}"
    wage[k] = float(r[3]) * CURNAME[r[4]]; wage_src[k] = 'ind'

# v132: la serie mensual original tiene escalones artificiales (mensualizaciones año por año, empalmes con
# crecimiento nominal constante, trimestres planos). Se reconstruye la trayectoria mensual con un método de
# benchmarking tipo Denton: se minimizan los cambios mes a mes del salario REAL (deflactado por IPC) sujeto a
# que el PROMEDIO NOMINAL de cada año sea exactamente el de la serie original. En los años derivados de un
# índice mensual oficial se conserva el perfil mensual real original (sólo se suavizan los empalmes entre años).
import numpy as np
_types = {}
for r in list(_wb['Serie mensual 1940-1993'].iter_rows(values_only=True))[1:]:
    if r[1] and r[3] is not None: _types[f"{int(r[1])}-{int(r[2]):02d}"] = str(r[6] or '')
_K = sorted(k for k in wage if k in idx)
_n = len(_K)
_P = np.array([idx[k] for k in _K])
_q = np.array([wage[k] for k in _K]) / _P                 # salario real original
_off = np.array(['índice mensual oficial' in _types.get(k, '') for k in _K])
# Perfil mensual oficial suavizado con media móvil centrada de 5 meses (en log) para quitar picos estacionales
# (aguinaldo/vacaciones de enero, caída de febrero) que no son cambios de nivel del salario.
_lq = np.log(_q); _sm = _lq.copy()
for i in range(_n):
    w = [j for j in range(i - 2, i + 3) if 0 <= j < _n and _off[j] and _K[j][:4] == _K[i][:4]]
    if _off[i]: _sm[i] = np.mean(_lq[w])
_t = np.zeros(_n - 1)                                     # cambio real objetivo entre m-1 y m
for i in range(1, _n):
    if _off[i] and _off[i - 1] and _K[i][:4] == _K[i - 1][:4]: _t[i - 1] = (np.exp(_sm[i]) - np.exp(_sm[i - 1])) * np.exp(np.mean(_lq[[j for j in range(_n) if _K[j][:4] == _K[i][:4]]]) - np.mean(_sm[[j for j in range(_n) if _K[j][:4] == _K[i][:4]]]))
_D = np.zeros((_n - 1, _n)); _D[np.arange(_n - 1), np.arange(_n - 1)] = -1; _D[np.arange(_n - 1), np.arange(1, _n)] = 1
_years = sorted({k[:4] for k in _K})
_C = np.zeros((len(_years), _n)); _b = np.zeros(len(_years))
for j, y in enumerate(_years):
    m = np.array([k[:4] == y for k in _K]); _C[j, m] = _P[m] / m.sum(); _b[j] = (np.array([wage[k] for k in _K])[m]).mean()
_sc = 1 / np.median(_q)                                   # escala numérica
_A = np.block([[2 * _D.T @ _D, _C.T], [_C, np.zeros((len(_years), len(_years)))]])
_rhs = np.concatenate([2 * _D.T @ (_t * _sc), _b * _sc])
_r = np.linalg.solve(_A, _rhs)[:_n] / _sc
for k, rv, p in zip(_K, _r, _P): wage[k] = float(rv * p)
for y in _years:
    _orig = _b[_years.index(y)]; _new = np.mean([wage[k] for k in _K if k[:4] == y])
    assert abs(_new / _orig - 1) < 1e-9, (y, _orig, _new)
for k, v in ripte.items():
    if k >= '1994-07': wage[k] = v; wage_src[k] = 'ripte'
_a, _b = '1993-12', '1994-07'
_ra, _rb = wage[_a] / idx[_a], wage[_b] / idx[_b]
_gap = months(_a, _b)
for i, k in enumerate(_gap[1:-1], start=1):
    w = i / (len(_gap) - 1); wage[k] = (_ra ** (1 - w) * _rb ** w) * idx[k]; wage_src[k] = 'interp'

# ---------- IPC EE.UU. ----------
us = {}
with open(os.path.join(W, 'uscpi.csv')) as f:
    for r in csv.DictReader(f):
        try: us[r['observation_date'][:7]] = float(r['CPIAUCNS'])
        except ValueError: pass

# Oct-2025 no se publicó (cierre del gobierno federal de EE.UU.): se interpola geométricamente entre meses vecinos.
_um = months(min(us), max(us))
for i in range(1, len(_um) - 1):
    k = _um[i]
    if k not in us and _um[i-1] in us and _um[i+1] in us: us[k] = (us[_um[i-1]] * us[_um[i+1]]) ** 0.5
# ---------- Arrays mensuales ----------
END = max(last_cpi, max(fx), max(ripte))
AM = months(START, END)
def r6(x):
    if x is None: return None
    return float(f'{x:.6g}')
arr = lambda d: [r6(d.get(k)) for k in AM]
fx_free = {k: v for k, v in free_m.items() if k >= '1991-04'}

# ---------- PIB (variación anual real) ----------
gdp = {}
lv = {int(a[:4]): b for a, b in S['gdp1935']}
for y in range(1936, 1961):
    if y in lv and y - 1 in lv: gdp[y] = round((lv[y] / lv[y - 1] - 1) * 100, 1)
for y, v in S['wbGdp']:
    if v is not None and int(y) >= 1961: gdp[int(y)] = round(v, 1)

# ---------- Personas ----------
wp = json.load(open(os.path.join(W, 'wp.json')))
def person(x, kind):
    key = x.get('wiki') or x['name']; w = wp.get(key) or {}
    o = {'name': x['name'], 'start': x['start'], 'end': x['end']}
    if kind == 'p': o['type'] = x.get('type')
    else: o['ministry'] = x.get('ministry')
    if w and not w.get('missing'):
        o['wiki'] = w.get('url'); o['photo'] = (w.get('thumb') or '').split('?')[0] or None
    if x.get('note'): o['note'] = x['note']
    if 'Junta' in x['name']: o['photo'] = None
    return o
P = [person(x, 'p') for x in json.load(open(os.path.join(W, 'people', 'presidents.json')))]
Mn = [person(x, 'm') for x in json.load(open(os.path.join(W, 'people', 'ministers.json')))]

# ---------- Bienes ----------
goods = json.load(open(os.path.join(W, 'people', 'goods.json')))
staples = json.load(open(os.path.join(W, 'people', 'staples.json')))
CURF = {'m$n': 1e-13, '$ley': 1e-11, '$a': 1e-7, 'A': 1e-4, '$': 1.0}
def pts(g, usd=False):
    out = []
    for p in g.get('points', []):
        if p.get('doubtful'): continue
        o = {'p': p['period'] if len(p['period']) == 7 else p['period'] + '-07', 'src': p['source']}
        if p.get('model'): o['model'] = p['model']
        if usd: o['usd'] = p['price']
        else:
            if p['currency'] not in CURF: continue
            o['ars'] = r6(p['price'] * CURF[p['currency']]); o['cur'] = p['currency']; o['nominal'] = p['price']
        if p.get('note') and 'intervened' in p['note']: o['flag'] = 'INDEC intervenido (2007–2015): precios oficiales subestimados'
        out.append(o)
    return out
nat_prices = {'bread': mp(S['bread']), 'beef': mp(S['beef']), 'milk': mp(S['milk'])}
basket = {}
for key, label in [('bread', 'Pan francés (1 kg)'), ('beef', 'Asado (1 kg)'), ('milk', 'Leche entera (1 litro)')]:
    pp = pts(staples.get(key, {}))
    have = {p['p'] for p in pp}
    for k, v in nat_prices[key].items():   # serie mensual INDEC 2016-04+ (precios promedio GBA)
        if k not in have: pp.append({'p': k, 'ars': r6(v), 'cur': '$', 'nominal': v, 'src': 'https://www.indec.gob.ar/ftp/cuadros/economia/sh_ipc_precios_promedio.xls'})
    pp.sort(key=lambda x: x['p'])
    basket[key] = {'label': label, 'area': 'GBA', 'points': pp}
basket['car'] = {'label': 'Auto 0 km (el más barato con precio publicado)', 'kind': 'durable', 'points': pts(goods['car'])}
basket['apartment'] = {'label': 'Departamento de 50 m² en CABA', 'kind': 'durable', 'm2': 50, 'points': pts(goods['apartment'], usd=True),
                       'note': 'Departamento: precio de oferta en USD/m² (GCBA 2001–2011; IDECBA desde 2016; otras referencias en data/precios.xlsx).'}
for _k, _per in [('bread', 'kg'), ('beef', 'kg'), ('milk', 'litros')]: basket[_k]['kind'] = 'food'; basket[_k]['per'] = _per

# v129: base de precios históricos aportada (data/precios.xlsx, 1940-2026, precios nominales con su moneda).
MESES = {'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4, 'mayo': 5, 'junio': 6, 'julio': 7, 'agosto': 8, 'septiembre': 9, 'setiembre': 9, 'octubre': 10, 'noviembre': 11, 'diciembre': 12}
PRODUCTS = {  # (producto, unidad) -> (clave, etiqueta, tipo, unidad para "por sueldo")
    ('Pan', '1 kg'): ('bread', None, 'food', 'kg'), ('Asado', '1 kg'): ('beef', None, 'food', 'kg'), ('Leche', '1 L'): ('milk', None, 'food', 'litros'),
    ('Auto nuevo', '1 unidad'): ('car', None, 'durable', None), ('Departamento', '1 m²'): ('apartment', None, 'durable', None),
    ('Aceite', '1 kg'): ('oil', 'Aceite (1 kg)', 'food', 'kg'), ('Aceite', '1.5 L'): ('oil15', 'Aceite (botella de 1,5 L)', 'food', 'botellas'),
    ('Aceite', '1 L'): ('oil1', 'Aceite (1 litro)', 'food', 'litros'), ('Harina', '1 kg'): ('flour', 'Harina de trigo (1 kg)', 'food', 'kg'),
    ('Pescado', '1 kg'): ('fish', 'Pescado (corvina, 1 kg)', 'food', 'kg'), ('Papa', '1 kg'): ('potato', 'Papa (1 kg)', 'food', 'kg'),
    ('Pollo', '1 kg'): ('chicken', 'Pollo (1 kg)', 'food', 'kg'), ('Azúcar', '1 kg'): ('sugar', 'Azúcar (1 kg)', 'food', 'kg'),
    ('Manteca', '1 kg'): ('butter', 'Manteca (1 kg)', 'food', 'kg'), ('Vino', '1 L'): ('wine', 'Vino común (1 litro)', 'food', 'litros'),
    ('Nafta', '1 L'): ('fuel', 'Nafta (1 litro)', 'food', 'litros'), ('Huevos', '12 unidades'): ('eggs', 'Huevos (docena)', 'food', 'docenas'),
    ('Café', '250 g'): ('coffee', 'Café molido (250 g)', 'food', 'paquetes'), ('Yerba mate', '1 kg'): ('yerba', 'Yerba mate (1 kg)', 'food', 'kg'),
    ('Arroz', '1 kg'): ('rice', 'Arroz (1 kg)', 'food', 'kg'), ('Motocicleta nueva', '1 unidad'): ('moto', 'Moto nueva', 'durable', None),
}
def obs_month(r):
    year, tipo, model, notes = int(r[0]), str(r[7] or ''), str(r[5] or ''), str(r[10] or '')
    if len(r) > 11 and r[11]: return f'{year}-{int(r[11]):02d}'   # columna "Mes" (agregada en v130)
    if 'diciembre' in tipo.lower(): return f'{year}-12'
    if 'julio' in tipo.lower(): return f'{year}-07'
    for txt in (model, notes):
        for name, mm in MESES.items():
            if re.search(rf'\b{name}\b', txt, re.I): return f'{year}-{mm:02d}'
    return f'{year}-07'   # promedio anual u observación sin mes: mitad de año
_pw = openpyxl.load_workbook(os.path.join(ROOT, 'data', 'precios.xlsx'), data_only=True)
added, skipped = 0, []
for r in list(_pw['Base'].iter_rows(values_only=True))[1:]:
    if not r[0] or r[3] is None: continue
    spec = PRODUCTS.get((r[1], r[2]))
    if not spec: skipped.append(f'{r[0]} {r[1]} ({r[2]})'); continue
    key, label, kind, per = spec
    item = basket.setdefault(key, {'label': label, 'kind': kind, 'points': []})
    if per: item['per'] = per
    pm = obs_month(r)
    o = {'p': pm, 'src': r[9] or r[8], 'model': (r[5] or None), 'area': r[6], 'type': r[7]}
    fxm = fx.get(pm) or fx.get(min(pm, max(fx)))
    if key == 'apartment':                       # departamentos: siempre en USD/m²
        if r[4] == 'USD': o['usd'] = float(r[3])
        elif fxm: o['usd'] = float(r[3]) * CURNAME[r[4]] / fxm; o['nominalArs'] = float(r[3])
        else: continue
    elif r[4] == 'USD':                          # precio publicado en dólares (convertibilidad): a pesos con el tipo de cambio del mes
        if not fxm: continue
        o['ars'] = r6(float(r[3]) * fxm); o['cur'] = 'USD'; o['nominal'] = float(r[3])
    else: o['ars'] = r6(float(r[3]) * CURNAME[r[4]]); o['cur'] = r[4]; o['nominal'] = float(r[3])
    if 'regulado' in str(r[7]).lower(): o['flag'] = 'Precio máximo regulado'
    o = {k: v for k, v in o.items() if v not in (None, '')}
    have = {p['p']: p for p in item['points']}
    if pm in have:
        if key == 'car' and o.get('ars') and have[pm].get('ars') and o['ars'] < have[pm]['ars']: item['points'].remove(have[pm])  # auto: el más barato del período
        else: continue
    item['points'].append(o); added += 1

# Control de saltos: observaciones contiguas (<= 24 meses) cuyo precio REAL se multiplica o divide por más de 3.
outliers = []
for key, item in basket.items():
    keep = []
    for p in sorted(item['points'], key=lambda x: x['p']):
        q = keep[-1] if keep else None
        if key != 'car' and q and 'ars' in p and 'ars' in q and q['p'] in idx and min(p['p'], last_cpi) in idx and len(months(q['p'], p['p'])) <= 25:
            ratio = (p['ars'] / idx[min(p['p'], last_cpi)]) / (q['ars'] / idx[min(q['p'], last_cpi)])
            if ratio > 3 or ratio < 1 / 3: outliers.append(f"{key} {p['p']} ({p.get('nominal')} {p.get('cur')}; real x{ratio:.1f} vs {q['p']})"); continue
        keep.append(p)
    item['points'] = keep
print('precios agregados', added, '| omitidos', skipped, '| saltos excluidos', outliers)

out = {
    'schema': 1, 'generated': today, 'start': START, 'end': END, 'cpiBase': last_cpi,
    'notes': {
        'cpi': 'IPC INDEC (GBA) 1943–2006; IPC San Luis 2007–2016 (INDEC intervenido); IPC Nacional INDEC desde 2017.',
        'fx': 'Tipo de cambio oficial/de referencia BCRA (series históricas monetarias); desde mayo 2026, promedio mensual del oficial. Dólar libre: promedio mensual desde 1991 y referencia anual antes.',
        'salary': 'Salario: hasta 1993, salario medio industrial nominal mensual (reconstrucción con fuentes INDEC, CEPAL, BCRA y Ministerio de Economía); desde julio de 1994, RIPTE (sueldo promedio de trabajadores registrados estables). Enero-junio 1994: interpolado con IPC. La trayectoria mensual hasta 1993 se suaviza en términos reales (método Denton) respetando el promedio nominal de cada año de la serie original.',
        'prices': 'Precios: INDEC (precios promedio), Anuario Estadístico, GCBA/IDECBA, listas y avisos contemporáneos (fuente enlazada en cada precio). Entre dos observaciones separadas por hasta 2 años, el precio del mes se interpola manteniendo la trayectoria real (ajustada por IPC); sin observaciones cercanas no se muestra.',
        'gdp': 'PIB a precios constantes: Cuentas Nacionales 1935–1962 (Secretaría de Asuntos Económicos) hasta 1960; Banco Mundial (INDEC) desde 1961.',
    },
    'cpi': arr(cpi), 'usCpi': arr(us), 'fx': arr(fx), 'fxFree': arr(fx_free), 'wage': arr(wage), 'wageSrc': [wage_src.get(k) for k in AM], 'ripte': arr(ripte),
    'fxFreeAnnual': {k: r6(v) for k, v in fx_free_annual.items()},
    'gdp': gdp,
    'currencies': [
        {'from': '1881-11-05', 'to': '1969-12-31', 'name': 'Peso moneda nacional', 'sym': 'm$n', 'f': 1e-13},
        {'from': '1970-01-01', 'to': '1983-05-31', 'name': 'Peso ley 18.188', 'sym': '$ ley', 'f': 1e-11},
        {'from': '1983-06-01', 'to': '1985-06-14', 'name': 'Peso argentino', 'sym': '$a', 'f': 1e-7},
        {'from': '1985-06-15', 'to': '1991-12-31', 'name': 'Austral', 'sym': '₳', 'f': 1e-4},
        {'from': '1992-01-01', 'to': None, 'name': 'Peso', 'sym': '$', 'f': 1.0},
    ],
    'presidents': P, 'ministers': Mn, 'basket': basket,
}
dst = os.path.join(ROOT, 'public', 'birth-economy.json')
json.dump(out, open(dst, 'w'), ensure_ascii=False, separators=(',', ':'))
print('ok', dst, os.path.getsize(dst), 'meses', len(AM), AM[0], AM[-1], 'IPC base', last_cpi)
